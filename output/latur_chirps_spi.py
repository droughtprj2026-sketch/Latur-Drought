# -*- coding: utf-8 -*-
"""
=====================================================================
 Monthly SPI (Standardized Precipitation Index) from CHIRPS rainfall
 Latur district, Marathwada  |  2016 - 2026  |  Google Colab script
=====================================================================

HOW TO USE IN GOOGLE COLAB
--------------------------
1. Zip your folder  D:\\Droughtprj@2026\\output\\chirps\\  (with the year
   sub-folders 2016 ... 2026 inside) -> e.g. chirps.zip
2. In Colab, either
      a) upload chirps.zip when the script asks, or
      b) put chirps.zip in Google Drive and set ZIP_PATH below.
3. Run the whole script in one cell:   %run latur_chirps_spi.py
   (or paste it into a cell).

WHAT IT DOES
------------
* Unzips and finds every .tif, reading YEAR and MONTH from the FILE NAME
  (not the folder), so the misplaced  2024/Latur_precip_2026_08.tif  is
  placed correctly as Aug 2026.
* Handles your two naming styles:
      CHIRPS_Latur_YYYY_MM_mm.tif   and   Latur_precip_YYYY_MM.tif
  2021-11, 2021-12 and all of 2022 exist in BOTH styles -> only one is
  kept (see PREFERRED_PREFIX) and the two are compared and reported.
* Resamples every raster onto ONE common grid (the two styles may have
  slightly different extent/resolution).
* Computes SPI-1, SPI-3, SPI-6, SPI-12 per pixel using the gamma
  distribution (McKee et al., 1993), fitted separately for each
  calendar month, with zero-rainfall correction.

SPI FORMULA (implemented below in `spi_from_series`)
-----------------------------------------------------
For one pixel, one calendar month, accumulated rainfall x (n years):
  1. A     = ln(mean(x)) - mean(ln(x))          (non-zero x only)
  2. alpha = (1 + sqrt(1 + 4A/3)) / (4A)        (shape, Thom 1958)
  3. beta  = mean(x) / alpha                    (scale)
  4. G(x)  = Gamma CDF(x; alpha, beta)
  5. q     = (number of zeros) / n
     H(x)  = q + (1 - q) * G(x)     (zero months: H = q/2, see ZERO_CENTRE)
  6. SPI   = inverse standard normal of H(x)

OUTPUTS  (folder: OUTPUT_DIR, also copied to Google Drive if mounted)
-------
  geojson/SPI_Latur_YYYY.geojson    one per year; one polygon per pixel,
                                    columns SPI1_01..SPI1_12, SPI3_01.. etc.
  csv/SPI_Latur_YYYY.csv            one per year; pixel x month table
  csv/SPI_Latur_district_monthly.csv
  SPI_Latur_2016_2026_combined.xlsx all sheets combined
  maps/SPI{k}_maps_YYYY.png         12-month map panel per year
  charts/*.png                      time series, drought frequency
  heatmaps/*.png                    year x month heat maps
  rasters/SPI{k}/SPI{k}_YYYY_MM.tif monthly SPI GeoTIFFs

NOTE ON RECORD LENGTH
---------------------
Only ~10 years are available. WMO recommends >= 30 years of record for
SPI. With 10 values per calendar month the gamma fit is rough, so treat
values beyond +/-2 with caution. Set CALIB_START / CALIB_END when you add
older CHIRPS data (available from 1981).
"""

# =====================================================================
# 0. SETTINGS - change these if needed
# =====================================================================
ZIP_PATH = None                 # e.g. '/content/drive/MyDrive/chirps.zip'; None = upload prompt
DATA_DIR = '/content/chirps_data'
OUTPUT_DIR = '/content/SPI_output'
DRIVE_COPY_DIR = '/content/drive/MyDrive/Droughtprj2026/SPI_output'   # None = don't copy

SPI_SCALES = [1, 3, 6, 12]      # months
PREFERRED_PREFIX = 'CHIRPS_'    # when a month exists twice, keep this style
                                # (set to 'Latur_precip' to prefer the other)
CALIB_START = None              # e.g. 2016 ; None = use all years
CALIB_END = None                # e.g. 2025
ZERO_CENTRE = True             # H(0)=q/2 (recommended) instead of H(0)=q
MIN_NONZERO = 3                 # min non-zero values needed for a gamma fit
NODATA_OUT = -9999.0

DISTRICT_NAME = 'Latur'

# =====================================================================
# 1. IMPORTS (installs missing packages in Colab)
# =====================================================================
import os, re, glob, shutil, zipfile, warnings, logging, subprocess, sys

def _ensure(pkgs):
    for imp, pipname in pkgs:
        try:
            __import__(imp)
        except ImportError:
            subprocess.check_call([sys.executable, '-m', 'pip', 'install', '-q', pipname])

_ensure([('rasterio', 'rasterio'), ('geopandas', 'geopandas'),
         ('openpyxl', 'openpyxl'), ('shapely', 'shapely')])

import numpy as np
import pandas as pd
from scipy import stats
import matplotlib
import matplotlib.pyplot as plt
from matplotlib.colors import BoundaryNorm, ListedColormap
from matplotlib.patches import Patch

warnings.filterwarnings('ignore')
logging.getLogger('rasterio._env').setLevel(logging.ERROR)
logging.getLogger('rasterio').setLevel(logging.ERROR)

# =====================================================================
# 2. DROUGHT CLASSES (McKee et al., 1993)
# =====================================================================
SPI_BOUNDS = [-10, -2.0, -1.5, -1.0, 1.0, 1.5, 2.0, 10]
SPI_LABELS = ['Extremely dry', 'Severely dry', 'Moderately dry',
              'Near normal', 'Moderately wet', 'Very wet', 'Extremely wet']
SPI_COLORS = ['#7f0000', '#d7301f', '#fdae61', '#f7f7f7',
              '#a6d96a', '#1a9850', '#00441b']
SPI_CMAP = ListedColormap(SPI_COLORS)
SPI_NORM = BoundaryNorm(SPI_BOUNDS, SPI_CMAP.N)
MONTHS = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun',
          'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']


def classify_spi(v):
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return 'No data'
    for lo, hi, lab in zip(SPI_BOUNDS[:-1], SPI_BOUNDS[1:], SPI_LABELS):
        if lo <= v < hi:
            return lab
    return SPI_LABELS[-1]


# =====================================================================
# 3. GET THE DATA (unzip)
# =====================================================================
def get_data():
    os.makedirs(DATA_DIR, exist_ok=True)
    zp = ZIP_PATH
    if zp is None:
        try:
            from google.colab import files
            print('Upload your CHIRPS zip file ...')
            up = files.upload()
            zp = list(up.keys())[0]
        except ImportError:
            raise SystemExit('Not in Colab: set ZIP_PATH to your zip file.')
    print('Extracting', zp)
    with zipfile.ZipFile(zp) as z:
        z.extractall(DATA_DIR)


# =====================================================================
# 4. FIND FILES, PARSE DATES FROM NAMES, REMOVE DUPLICATES
# =====================================================================
NAME_RE = re.compile(r'(19|20)(\d{2})[_\-](\d{2})')


def list_files():
    tifs = sorted(glob.glob(os.path.join(DATA_DIR, '**', '*.tif'), recursive=True))
    tifs = [t for t in tifs if not os.path.basename(t).startswith('._')]  # mac junk
    rows = []
    for t in tifs:
        name = os.path.basename(t)
        m = NAME_RE.search(name)
        if not m:
            print('  skipped (no date in name):', name)
            continue
        year = int(m.group(1) + m.group(2))
        month = int(m.group(3))
        if not 1 <= month <= 12:
            print('  skipped (bad month):', name)
            continue
        rows.append(dict(path=t, file=name, year=year, month=month,
                         naming='CHIRPS_' if name.startswith('CHIRPS_') else 'Latur_precip',
                         folder=os.path.basename(os.path.dirname(t))))
    df = pd.DataFrame(rows)
    if df.empty:
        raise SystemExit('No .tif files with YYYY_MM in their names were found.')

    moved = df[df['folder'].str.isdigit() & (df['folder'] != df['year'].astype(str))]
    for _, r in moved.iterrows():
        print(f"  note: {r.file} is in folder {r.folder} -> used as {r.year}-{r.month:02d}")

    df['pref'] = (~df['file'].str.startswith(PREFERRED_PREFIX)).astype(int)
    dup_keys = df[df.duplicated(['year', 'month'], keep=False)]
    df = df.sort_values(['year', 'month', 'pref']).reset_index(drop=True)
    kept = df.drop_duplicates(['year', 'month'], keep='first').reset_index(drop=True)
    return kept, dup_keys


# =====================================================================
# 5. READ ALL RASTERS ONTO ONE COMMON GRID
# =====================================================================
def read_stack(files_df):
    import rasterio
    from rasterio.warp import reproject, Resampling

    # reference grid = finest-resolution file (so nothing is down-sampled)
    best = None
    for p in files_df['path']:
        with rasterio.open(p) as src:
            res = abs(src.transform.a)
            if best is None or res < best[0] - 1e-12:
                best = (res, p)
    with rasterio.open(best[1]) as ref:
        profile = ref.profile.copy()
        H, W = ref.height, ref.width
        ref_transform, ref_crs = ref.transform, ref.crs
    print(f'Reference grid: {os.path.basename(best[1])}  ({H} x {W} px, res {best[0]:.5f})')

    # complete monthly calendar (gaps become NaN)
    start = pd.Timestamp(int(files_df.year.min()), 1, 1)
    last = files_df.sort_values(['year', 'month']).iloc[-1]
    end = pd.Timestamp(int(last.year), int(last.month), 1)
    dates = pd.date_range(start, end, freq='MS')
    lut = {(r.year, r.month): r.path for r in files_df.itertuples()}

    stack = np.full((len(dates), H, W), np.nan, dtype='float64')
    missing = []
    for i, d in enumerate(dates):
        p = lut.get((d.year, d.month))
        if p is None:
            missing.append(d.strftime('%Y-%m'))
            continue
        with rasterio.open(p) as src:
            arr = src.read(1).astype('float64')
            nd = src.nodata
            if nd is not None:
                arr[arr == nd] = np.nan
            arr[(arr < 0) | (arr > 5000)] = np.nan       # -9999 and junk
            same = (src.height == H and src.width == W and
                    src.crs == ref_crs and
                    np.allclose(tuple(src.transform)[:6], tuple(ref_transform)[:6], atol=1e-9))
            if same:
                stack[i] = arr
            else:
                dst = np.full((H, W), np.nan)
                reproject(arr, dst, src_transform=src.transform, src_crs=src.crs,
                          src_nodata=np.nan, dst_transform=ref_transform,
                          dst_crs=ref_crs, dst_nodata=np.nan,
                          resampling=Resampling.bilinear)
                stack[i] = dst
    if missing:
        print('  missing months (left as NaN):', ', '.join(missing))

    # pixels outside the district = NaN in every month
    valid_mask = np.isfinite(stack).any(axis=0)
    profile.update(dtype='float32', count=1, nodata=NODATA_OUT, compress='lzw')
    return stack, dates, profile, ref_transform, ref_crs, valid_mask


def compare_duplicates(dup_df, profile_transform=None):
    """Print mean rainfall of both naming styles where a month exists twice."""
    if dup_df.empty:
        return pd.DataFrame()
    import rasterio
    rows = []
    for (y, m), g in dup_df.groupby(['year', 'month']):
        rec = {'year': y, 'month': m}
        for r in g.itertuples():
            with rasterio.open(r.path) as src:
                a = src.read(1).astype(float)
                if src.nodata is not None:
                    a[a == src.nodata] = np.nan
                a[a < 0] = np.nan
                rec[f'{r.naming}_mean_mm'] = round(float(np.nanmean(a)), 2)
                rec[f'{r.naming}_shape'] = f'{src.height}x{src.width}'
        rows.append(rec)
    cmp_df = pd.DataFrame(rows)
    print('\nMonths found in BOTH naming styles (mean rainfall, mm):')
    print(cmp_df.to_string(index=False))
    return cmp_df


# =====================================================================
# 6. SPI CORE - the formula
# =====================================================================
def rolling_sum(stack, k):
    """k-month accumulation along time axis; first k-1 months = NaN."""
    if k == 1:
        return stack.copy()
    T = stack.shape[0]
    out = np.full_like(stack, np.nan)
    for t in range(k - 1, T):
        win = stack[t - k + 1:t + 1]
        ok = np.isfinite(win).all(axis=0)          # need all k months present
        s = np.nansum(win, axis=0)
        s[~ok] = np.nan
        out[t] = s
    return out


def gamma_params(x):
    """Thom (1958) maximum-likelihood estimates. x = non-zero values."""
    mean = x.mean()
    A = np.log(mean) - np.mean(np.log(x))
    if A <= 0:                                    # all values identical
        return np.nan, np.nan
    alpha = (1.0 + np.sqrt(1.0 + 4.0 * A / 3.0)) / (4.0 * A)
    beta = mean / alpha
    return alpha, beta


def spi_from_series(x_all, calib_idx=None):
    """
    x_all : 1-D array of accumulated rain for ONE pixel, ONE calendar month
            (one value per year), may contain NaN.
    returns SPI array, same length.
    """
    out = np.full(x_all.shape, np.nan)
    good = np.isfinite(x_all)
    if calib_idx is None:
        calib = x_all[good]
    else:
        c = x_all[calib_idx]
        calib = c[np.isfinite(c)]
    n = calib.size
    if n == 0:
        return out
    nonzero = calib[calib > 0]
    q = (n - nonzero.size) / n                   # probability of zero rain

    if nonzero.size < MIN_NONZERO:
        return out
    alpha, beta = gamma_params(nonzero)
    if not np.isfinite(alpha):
        return out

    x = x_all[good]
    G = stats.gamma.cdf(x, a=alpha, scale=beta)
    Hx = q + (1.0 - q) * G
    # H(0) = q in the classic formula. In dry months (Jan-May) q can be large,
    # and H(0)=q would make every zero-rain month look "wet". Using the centre
    # of the zero-probability mass, q/2 (Stagge et al., 2015), fixes this.
    # Set ZERO_CENTRE = False for the classic H(0) = q.
    Hx = np.where(x <= 0, q / 2.0 if ZERO_CENTRE else q, Hx)
    Hx = np.clip(Hx, 0.0001, 0.9999)
    out[good] = stats.norm.ppf(Hx)
    return out


def compute_spi(stack, dates, k, valid_mask):
    acc = rolling_sum(stack, k)
    T, H, W = acc.shape
    spi = np.full_like(acc, np.nan)
    years = dates.year.values
    for m in range(1, 13):
        tidx = np.where(dates.month == m)[0]
        if tidx.size == 0:
            continue
        calib_idx = None
        if CALIB_START or CALIB_END:
            ys = years[tidx]
            lo = CALIB_START or ys.min()
            hi = CALIB_END or ys.max()
            calib_idx = (ys >= lo) & (ys <= hi)
        sub = acc[tidx]                           # (n_years, H, W)
        for r in range(H):
            for c in range(W):
                if valid_mask[r, c]:
                    spi[tidx, r, c] = spi_from_series(sub[:, r, c], calib_idx)
    return spi


# =====================================================================
# 7. OUTPUT WRITERS
# =====================================================================
def save_rasters(spi_all, dates, profile):
    import rasterio
    for k, spi in spi_all.items():
        d = os.path.join(OUTPUT_DIR, 'rasters', f'SPI{k}')
        os.makedirs(d, exist_ok=True)
        for i, dt in enumerate(dates):
            a = spi[i].astype('float32')
            if not np.isfinite(a).any():
                continue
            a = np.where(np.isfinite(a), a, NODATA_OUT).astype('float32')
            with rasterio.open(os.path.join(d, f'SPI{k}_{dt:%Y_%m}.tif'), 'w', **profile) as dst:
                dst.write(a, 1)


def pixel_table(spi_all, dates, transform, valid_mask, year):
    """One row per valid pixel; columns SPI{k}_{MM} for the given year."""
    rr, cc = np.where(valid_mask)
    xs = transform.c + (cc + 0.5) * transform.a
    ys = transform.f + (rr + 0.5) * transform.e
    df = pd.DataFrame({'pixel_id': np.arange(rr.size), 'row': rr, 'col': cc,
                       'lon': np.round(xs, 6), 'lat': np.round(ys, 6)})
    tidx = np.where(dates.year == year)[0]
    for k, spi in spi_all.items():
        for i in tidx:
            m = dates[i].month
            df[f'SPI{k}_{m:02d}'] = np.round(spi[i][rr, cc], 3)
    # yearly summary per pixel (SPI-1)
    s1 = [c for c in df.columns if c.startswith('SPI1_')]
    if s1:
        df['SPI1_min'] = df[s1].min(axis=1)
        df['SPI1_mean'] = df[s1].mean(axis=1).round(3)
        df['dry_months'] = (df[s1] <= -1.0).sum(axis=1)
    return df, rr, cc


def save_geojson_csv(spi_all, dates, transform, crs, valid_mask):
    import geopandas as gpd
    from shapely.geometry import box
    gdir = os.path.join(OUTPUT_DIR, 'geojson'); os.makedirs(gdir, exist_ok=True)
    cdir = os.path.join(OUTPUT_DIR, 'csv'); os.makedirs(cdir, exist_ok=True)
    px_tables = {}
    for y in sorted(set(dates.year)):
        df, rr, cc = pixel_table(spi_all, dates, transform, valid_mask, y)
        df.to_csv(os.path.join(cdir, f'SPI_Latur_{y}.csv'), index=False)
        geoms = []
        for r, c in zip(rr, cc):
            x0 = transform.c + c * transform.a
            y0 = transform.f + r * transform.e
            geoms.append(box(min(x0, x0 + transform.a), min(y0, y0 + transform.e),
                             max(x0, x0 + transform.a), max(y0, y0 + transform.e)))
        gdf = gpd.GeoDataFrame(df, geometry=geoms, crs=crs)
        if gdf.crs is not None and gdf.crs.to_epsg() != 4326:
            gdf = gdf.to_crs(4326)
        gdf = gdf.where(pd.notna(gdf), None)
        gdf.to_file(os.path.join(gdir, f'SPI_Latur_{y}.geojson'), driver='GeoJSON')
        px_tables[y] = df
        print(f'  {y}: {len(df)} pixels -> geojson + csv')
    return px_tables


def district_table(stack, spi_all, dates, valid_mask):
    rec = pd.DataFrame({'date': dates.strftime('%Y-%m'),
                        'year': dates.year, 'month': dates.month,
                        'month_name': [MONTHS[m - 1] for m in dates.month]})
    rain = stack[:, valid_mask]
    rec['rain_mm_mean'] = np.round(np.nanmean(rain, axis=1), 2)
    for k, spi in spi_all.items():
        v = spi[:, valid_mask]
        mean = np.nanmean(v, axis=1)
        rec[f'SPI{k}'] = np.round(mean, 3)
        rec[f'SPI{k}_class'] = [classify_spi(x) for x in mean]
        with np.errstate(invalid='ignore'):
            n = np.isfinite(v).sum(axis=1)
            rec[f'SPI{k}_pct_area_dry'] = np.where(
                n > 0, np.round(100 * (v <= -1.0).sum(axis=1) / np.maximum(n, 1), 1), np.nan)
    return rec


def yearly_summary(dist):
    rows = []
    for y, g in dist.groupby('year'):
        r = {'year': y, 'annual_rain_mm': round(g['rain_mm_mean'].sum(), 1),
             'monsoon_JJAS_rain_mm': round(g[g.month.between(6, 9)]['rain_mm_mean'].sum(), 1),
             'months_available': int(g['rain_mm_mean'].notna().sum())}
        for k in SPI_SCALES:
            col = f'SPI{k}'
            if col in g:
                r[f'{col}_mean'] = round(g[col].mean(), 3)
                r[f'{col}_min'] = round(g[col].min(), 3)
                r[f'{col}_dry_months(<=-1)'] = int((g[col] <= -1).sum())
        if 'SPI12' in g and (g.month == 12).any():
            dec = g.loc[g.month == 12, 'SPI12'].values[0]
            r['SPI12_Dec'] = dec
            r['SPI12_Dec_class'] = classify_spi(dec)
        if 'SPI3' in g and (g.month == 9).any():
            sep = g.loc[g.month == 9, 'SPI3'].values[0]       # JAS monsoon
            r['SPI3_Sep(Jul-Sep)'] = sep
            r['SPI3_Sep_class'] = classify_spi(sep)
        rows.append(r)
    return pd.DataFrame(rows)


def class_counts(dist):
    out = []
    for k in SPI_SCALES:
        col = f'SPI{k}_class'
        if col not in dist:
            continue
        vc = dist[col].value_counts()
        out.append({'scale': f'SPI-{k}', **{lab: int(vc.get(lab, 0)) for lab in SPI_LABELS}})
    return pd.DataFrame(out)


def save_excel(dist, ysum, ccounts, px_tables, files_df, dup_cmp, path):
    with pd.ExcelWriter(path, engine='openpyxl') as xw:
        readme = pd.DataFrame({'Item': [
            'Study area', 'Data', 'Period', 'Method', 'Timescales',
            'Drought classes', 'Record-length note'],
            'Description': [
            f'{DISTRICT_NAME} district, Marathwada, Maharashtra',
            'CHIRPS monthly precipitation (mm)',
            f"{dist['date'].iloc[0]} to {dist['date'].iloc[-1]}",
            'Gamma distribution fitted per pixel and calendar month (Thom MLE), '
            'zero-rain correction H=q+(1-q)G, SPI = inverse normal of H (McKee 1993)',
            ', '.join(f'SPI-{k}' for k in SPI_SCALES),
            '<=-2 Extreme dry | -2..-1.5 Severe | -1.5..-1 Moderate | -1..1 Normal | '
            '1..1.5 Moderately wet | 1.5..2 Very wet | >=2 Extremely wet',
            '~10 years only; WMO recommends >=30 years. Interpret extremes with care.']})
        readme.to_excel(xw, sheet_name='README', index=False)
        dist.to_excel(xw, sheet_name='District_Monthly_SPI', index=False)
        ysum.to_excel(xw, sheet_name='Yearly_Summary', index=False)
        ccounts.to_excel(xw, sheet_name='Drought_Class_Counts', index=False)
        for k in SPI_SCALES:
            if f'SPI{k}' in dist:
                pv = dist.pivot(index='year', columns='month', values=f'SPI{k}')
                pv.columns = [MONTHS[m - 1] for m in pv.columns]
                pv.to_excel(xw, sheet_name=f'SPI{k}_YearxMonth')
        pv = dist.pivot(index='year', columns='month', values='rain_mm_mean')
        pv.columns = [MONTHS[m - 1] for m in pv.columns]
        pv.to_excel(xw, sheet_name='Rain_YearxMonth')
        for y, df in px_tables.items():
            df.to_excel(xw, sheet_name=f'Pixels_{y}', index=False)
        files_df[['year', 'month', 'file', 'folder']].to_excel(xw, sheet_name='Files_Used', index=False)
        if dup_cmp is not None and not dup_cmp.empty:
            dup_cmp.to_excel(xw, sheet_name='Duplicate_Check', index=False)

        # light formatting: widths, header fill, SPI colouring
        from openpyxl.styles import PatternFill, Font
        from openpyxl.formatting.rule import CellIsRule
        hdr = PatternFill('solid', fgColor='1F4E78')
        for ws in xw.book.worksheets:
            for cell in ws[1]:
                cell.fill = hdr; cell.font = Font(color='FFFFFF', bold=True)
            for col in ws.columns:
                w = max(len(str(c.value)) if c.value is not None else 0 for c in col[:200])
                ws.column_dimensions[col[0].column_letter].width = min(max(10, w + 2), 60)
            ws.freeze_panes = 'B2'
            if ws.title.startswith('SPI') and ws.title.endswith('YearxMonth'):
                rng = f'B2:M{ws.max_row}'
                rules = [('lessThanOrEqual', ['-2'], 'C00000'),
                         ('between', ['-1.999', '-1.5'], 'FF6B6B'),
                         ('between', ['-1.499', '-1'], 'FFC000'),
                         ('between', ['1', '1.499'], 'C6EFCE'),
                         ('between', ['1.5', '1.999'], '70AD47'),
                         ('greaterThanOrEqual', ['2'], '375623')]
                for op, f, color in rules:
                    ws.conditional_formatting.add(
                        rng, CellIsRule(operator=op, formula=f,
                                        fill=PatternFill('solid', fgColor=color)))


# =====================================================================
# 8. MAPS, CHARTS, HEAT MAPS (matplotlib)
# =====================================================================
def _legend_handles():
    return [Patch(facecolor=c, edgecolor='grey', label=l)
            for c, l in zip(SPI_COLORS, SPI_LABELS)]


def plot_year_maps(spi_all, dates, transform, k=1):
    d = os.path.join(OUTPUT_DIR, 'maps'); os.makedirs(d, exist_ok=True)
    spi = spi_all[k]
    H, W = spi.shape[1:]
    ext = [transform.c, transform.c + W * transform.a,
           transform.f + H * transform.e, transform.f]
    for y in sorted(set(dates.year)):
        fig, axes = plt.subplots(3, 4, figsize=(15, 11))
        for m in range(1, 13):
            ax = axes.flat[m - 1]
            idx = np.where((dates.year == y) & (dates.month == m))[0]
            if idx.size and np.isfinite(spi[idx[0]]).any():
                a = spi[idx[0]]
                ax.imshow(np.ma.masked_invalid(a), cmap=SPI_CMAP, norm=SPI_NORM,
                          extent=ext, interpolation='nearest')
                ax.set_title(f'{MONTHS[m-1]} {y}   mean {np.nanmean(a):+.2f}', fontsize=10)
            else:
                ax.text(0.5, 0.5, 'no data', ha='center', va='center',
                        transform=ax.transAxes, color='grey')
                ax.set_title(f'{MONTHS[m-1]} {y}', fontsize=10)
            ax.tick_params(labelsize=7)
        fig.suptitle(f'SPI-{k}  |  {DISTRICT_NAME} district  |  {y}', fontsize=15, weight='bold')
        fig.legend(handles=_legend_handles(), loc='lower center', ncol=7, fontsize=9, frameon=False)
        fig.tight_layout(rect=[0, 0.05, 1, 0.96])
        fig.savefig(os.path.join(d, f'SPI{k}_maps_{y}.png'), dpi=150)
        plt.close(fig)


def plot_charts(dist):
    d = os.path.join(OUTPUT_DIR, 'charts'); os.makedirs(d, exist_ok=True)
    x = pd.to_datetime(dist['date'])

    # (a) bar time series for every scale
    scales = [k for k in SPI_SCALES if f'SPI{k}' in dist]
    fig, axes = plt.subplots(len(scales), 1, figsize=(15, 3 * len(scales)), sharex=True)
    axes = np.atleast_1d(axes)
    for ax, k in zip(axes, scales):
        v = dist[f'SPI{k}'].values
        cols = [SPI_COLORS[np.searchsorted(SPI_BOUNDS, val, side='right') - 1]
                if np.isfinite(val) else 'white' for val in v]
        ax.bar(x, v, width=25, color=cols, edgecolor='grey', linewidth=0.3)
        for t, ls in [(-1, '--'), (-1.5, '--'), (-2, '-')]:
            ax.axhline(t, color='#7f0000', ls=ls, lw=0.7, alpha=0.6)
        ax.axhline(0, color='black', lw=0.6)
        ax.set_ylabel(f'SPI-{k}'); ax.set_ylim(-3.2, 3.2); ax.grid(axis='y', alpha=0.3)
    axes[0].set_title(f'District-mean SPI  |  {DISTRICT_NAME}  |  '
                      f"{dist['date'].iloc[0]} to {dist['date'].iloc[-1]}", weight='bold')
    fig.legend(handles=_legend_handles(), loc='lower center', ncol=7, fontsize=9, frameon=False)
    fig.tight_layout(rect=[0, 0.04, 1, 1])
    fig.savefig(os.path.join(d, 'SPI_timeseries_all_scales.png'), dpi=150); plt.close(fig)

    # (b) rainfall + SPI-3 line
    fig, ax1 = plt.subplots(figsize=(15, 5))
    ax1.bar(x, dist['rain_mm_mean'], width=25, color='#4a90d9', alpha=0.6, label='Rainfall (mm)')
    ax1.set_ylabel('Monthly rainfall (mm)', color='#2c5f8a')
    ax2 = ax1.twinx()
    if 'SPI3' in dist:
        ax2.plot(x, dist['SPI3'], color='#b2182b', lw=1.6, label='SPI-3')
    if 'SPI12' in dist:
        ax2.plot(x, dist['SPI12'], color='black', lw=1.4, ls='--', label='SPI-12')
    ax2.axhline(-1, color='#b2182b', lw=0.7, ls=':'); ax2.axhline(0, color='grey', lw=0.5)
    ax2.set_ylabel('SPI'); ax2.set_ylim(-3.2, 3.2)
    h1, l1 = ax1.get_legend_handles_labels(); h2, l2 = ax2.get_legend_handles_labels()
    ax1.legend(h1 + h2, l1 + l2, loc='upper left', frameon=False)
    ax1.set_title(f'Monthly rainfall and SPI  |  {DISTRICT_NAME}', weight='bold')
    fig.tight_layout(); fig.savefig(os.path.join(d, 'Rainfall_vs_SPI.png'), dpi=150); plt.close(fig)

    # (c) % area in drought (SPI-3 <= -1)
    if 'SPI3_pct_area_dry' in dist:
        fig, ax = plt.subplots(figsize=(15, 4))
        ax.fill_between(x, dist['SPI3_pct_area_dry'].fillna(0), color='#d7301f', alpha=0.7, step='mid')
        ax.set_ylabel('% of district area'); ax.set_ylim(0, 100)
        ax.set_title('Area under drought (SPI-3 ≤ −1)', weight='bold'); ax.grid(alpha=0.3)
        fig.tight_layout(); fig.savefig(os.path.join(d, 'Drought_area_percent_SPI3.png'), dpi=150); plt.close(fig)

    # (d) drought-month counts per year
    rows = []
    for y, g in dist.groupby('year'):
        rows.append({'year': y, **{lab: int((g['SPI3_class'] == lab).sum())
                                   for lab in SPI_LABELS[:3]}})
    cnt = pd.DataFrame(rows).set_index('year')
    ax = cnt.plot(kind='bar', stacked=True, figsize=(11, 4.5),
                  color=SPI_COLORS[:3], edgecolor='grey')
    ax.set_ylabel('Number of months'); ax.set_xlabel('')
    ax.set_title('Drought months per year (SPI-3)', weight='bold')
    ax.legend(frameon=False); plt.xticks(rotation=0); plt.tight_layout()
    plt.savefig(os.path.join(d, 'Drought_months_per_year_SPI3.png'), dpi=150); plt.close()


def plot_heatmaps(dist):
    d = os.path.join(OUTPUT_DIR, 'heatmaps'); os.makedirs(d, exist_ok=True)
    scales = [k for k in SPI_SCALES if f'SPI{k}' in dist]

    def _one(ax, pv, title, cmap, norm=None, fmt='{:+.2f}', vmin=None, vmax=None):
        im = ax.imshow(pv.values, cmap=cmap, norm=norm, vmin=vmin, vmax=vmax, aspect='auto')
        ax.set_xticks(range(12)); ax.set_xticklabels(MONTHS)
        ax.set_yticks(range(len(pv.index))); ax.set_yticklabels(pv.index)
        for i in range(pv.shape[0]):
            for j in range(pv.shape[1]):
                v = pv.values[i, j]
                if np.isfinite(v):
                    dark = (norm is not None and abs(v) >= 1.5)
                    ax.text(j, i, fmt.format(v), ha='center', va='center', fontsize=8,
                            color='white' if dark else 'black')
        ax.set_title(title, weight='bold')
        return im

    for k in scales:
        pv = dist.pivot(index='year', columns='month', values=f'SPI{k}').reindex(columns=range(1, 13))
        fig, ax = plt.subplots(figsize=(13, 0.55 * len(pv) + 2))
        _one(ax, pv, f'SPI-{k} heat map  |  {DISTRICT_NAME} (district mean)', SPI_CMAP, SPI_NORM)
        fig.legend(handles=_legend_handles(), loc='lower center', ncol=7, fontsize=8, frameon=False)
        fig.tight_layout(rect=[0, 0.07, 1, 1])
        fig.savefig(os.path.join(d, f'Heatmap_SPI{k}.png'), dpi=150); plt.close(fig)

    # combined panel
    fig, axes = plt.subplots(1, len(scales), figsize=(7 * len(scales), 0.5 * dist.year.nunique() + 2.5))
    for ax, k in zip(np.atleast_1d(axes), scales):
        pv = dist.pivot(index='year', columns='month', values=f'SPI{k}').reindex(columns=range(1, 13))
        _one(ax, pv, f'SPI-{k}', SPI_CMAP, SPI_NORM, fmt='{:+.1f}')
    fig.legend(handles=_legend_handles(), loc='lower center', ncol=7, fontsize=9, frameon=False)
    fig.suptitle(f'SPI heat maps  |  {DISTRICT_NAME}', weight='bold', fontsize=14)
    fig.tight_layout(rect=[0, 0.06, 1, 0.95])
    fig.savefig(os.path.join(d, 'Heatmap_all_scales.png'), dpi=150); plt.close(fig)

    # rainfall heat map
    pv = dist.pivot(index='year', columns='month', values='rain_mm_mean').reindex(columns=range(1, 13))
    fig, ax = plt.subplots(figsize=(13, 0.55 * len(pv) + 2))
    im = _one(ax, pv, f'Monthly rainfall (mm)  |  {DISTRICT_NAME}', 'Blues', fmt='{:.0f}')
    fig.colorbar(im, ax=ax, label='mm')
    fig.tight_layout(); fig.savefig(os.path.join(d, 'Heatmap_rainfall.png'), dpi=150); plt.close(fig)


# =====================================================================
# 9. MAIN
# =====================================================================
def main():
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    if not glob.glob(os.path.join(DATA_DIR, '**', '*.tif'), recursive=True):
        get_data()

    print('\n[1] Scanning files ...')
    files_df, dup_df = list_files()
    print(f'  {len(files_df)} unique months, {files_df.year.min()}-{files_df.year.max()}')
    dup_cmp = compare_duplicates(dup_df)

    print('\n[2] Reading rasters ...')
    stack, dates, profile, transform, crs, valid = read_stack(files_df)
    print(f'  stack: {stack.shape} (months, rows, cols), {valid.sum()} valid pixels')

    print('\n[3] Computing SPI ...')
    spi_all = {}
    for k in SPI_SCALES:
        spi_all[k] = compute_spi(stack, dates, k, valid)
        print(f'  SPI-{k} done  (range {np.nanmin(spi_all[k]):+.2f} to {np.nanmax(spi_all[k]):+.2f})')

    print('\n[4] Writing GeoTIFFs, GeoJSON, CSV ...')
    save_rasters(spi_all, dates, profile)
    px_tables = save_geojson_csv(spi_all, dates, transform, crs, valid)
    dist = district_table(stack, spi_all, dates, valid)
    dist.to_csv(os.path.join(OUTPUT_DIR, 'csv', 'SPI_Latur_district_monthly.csv'), index=False)
    ysum = yearly_summary(dist)
    ysum.to_csv(os.path.join(OUTPUT_DIR, 'csv', 'SPI_Latur_yearly_summary.csv'), index=False)
    ccounts = class_counts(dist)

    xlsx = os.path.join(OUTPUT_DIR, f'SPI_Latur_{dates[0].year}_{dates[-1].year}_combined.xlsx')
    save_excel(dist, ysum, ccounts, px_tables, files_df, dup_cmp, xlsx)
    print('  Excel:', xlsx)

    print('\n[5] Maps, charts, heat maps ...')
    for k in (1, 3):
        if k in spi_all:
            plot_year_maps(spi_all, dates, transform, k)
    plot_charts(dist)
    plot_heatmaps(dist)

    print('\n[6] Zipping outputs ...')
    zpath = shutil.make_archive(OUTPUT_DIR, 'zip', OUTPUT_DIR)
    print('  ', zpath)
    if DRIVE_COPY_DIR and os.path.isdir('/content/drive/MyDrive'):
        shutil.copytree(OUTPUT_DIR, DRIVE_COPY_DIR, dirs_exist_ok=True)
        shutil.copy2(zpath, os.path.dirname(DRIVE_COPY_DIR))
        print('  copied to Drive:', DRIVE_COPY_DIR)

    print('\nYearly summary:')
    print(ysum.to_string(index=False))

    try:                                           # show a preview in Colab
        from IPython.display import Image, display
        for p in ['heatmaps/Heatmap_all_scales.png', 'charts/SPI_timeseries_all_scales.png']:
            display(Image(os.path.join(OUTPUT_DIR, p)))
    except Exception:
        pass
    try:
        from google.colab import files
        files.download(zpath)
    except Exception:
        pass
    print('\nDone.')


if __name__ == '__main__':
    main()
